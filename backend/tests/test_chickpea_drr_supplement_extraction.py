"""Tests use synthetic workbook fixtures only, never disease observations."""

import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from extract_chickpea_drr_supplement import (
    FARMER_HEADERS,
    LITERATURE_HEADERS,
    extract,
)


class ChickpeaSupplementExtractionTests(unittest.TestCase):
    def make_fixture(self, path: Path) -> None:
        workbook = Workbook()
        literature = workbook.active
        literature.title = "Supplementary Table S1 DataFile"
        literature.cell(3, 1, "State")
        for col, header in enumerate(LITERATURE_HEADERS, start=1):
            literature.cell(3, col, header)
        for row in range(4, 101):
            literature.cell(row, 1, f"Fixture state {row}")
            literature.cell(row, 11, 12.5)

        farmer = workbook.create_sheet("Supplementary Figure S14")
        for col, header in enumerate(FARMER_HEADERS, start=1):
            farmer.cell(3, col, header)
        for row in range(4, 10):
            farmer.cell(row, 1, f"Field {row - 3}")
            farmer.cell(row, 5, "Fixture village" if row == 4 else f"Fixture village {row}")
            farmer.cell(row, 7, "Groundnut" if row == 4 else "Chickpea")
            farmer.cell(row, 10, ">40" if row == 4 else row)
        for col, header in enumerate([None, "Location", "DI (%)", "SMC (%)", "DI SEM", "SMC SEM"], start=1):
            farmer.cell(17, col, header)
        for row in range(18, 26):
            if row % 2 == 0:
                farmer.cell(row, 1, f"Fixture district {row}")
            farmer.cell(row, 2, "Fixture village" if row == 18 else f"Summary village {row}")
            farmer.cell(row, 3, 55)
            farmer.cell(row, 4, 10)
        workbook.save(path)

    def test_extracts_rows_without_losing_duplicate_headers_or_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "fixture.xlsx"
            output = root / "output"
            self.make_fixture(workbook_path)
            counts = extract(workbook_path, output)
            self.assertEqual(counts, {
                "literature_records": 97,
                "farmer_field_entries": 6,
                "farmer_field_summary_measurements": 8,
            })
            with (output / "literature_records.csv").open(encoding="utf-8", newline="") as stream:
                reader = csv.reader(stream)
                header = next(reader)
                first = next(reader)
            self.assertEqual(header.count("Variety"), 2)
            self.assertEqual(first[:2], ["Supplementary Table S1 DataFile", "4"])
            self.assertEqual(first[12], "12.5")
            with (output / "farmer_field_entries.csv").open(encoding="utf-8", newline="") as stream:
                reader = csv.reader(stream)
                field_header = next(reader)
                field_values = next(reader)
            field = dict(zip(field_header[:5], field_values[:5]))
            self.assertEqual(field["current_crop_identity_status"], "ambiguous_current_crop")
            self.assertEqual(field["previous_crop_flag"], "previous_crop_is_not_chickpea")
            self.assertEqual(field_header.count("Reading 1"), 2)
            first_reading_1 = field_header.index("Reading 1")
            self.assertEqual(field_values[first_reading_1], ">40")
            with (output / "farmer_field_summary_measurements.csv").open(
                encoding="utf-8", newline=""
            ) as stream:
                summary = next(csv.DictReader(stream))
            self.assertEqual(summary["candidate_related_field_excel_rows"], "4")
            self.assertEqual(summary["DI (%)"], "55")

    def test_refuses_overwrite_by_default(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "fixture.xlsx"
            output = root / "output"
            self.make_fixture(workbook_path)
            extract(workbook_path, output)
            with self.assertRaises(FileExistsError):
                extract(workbook_path, output)


if __name__ == "__main__":
    unittest.main()
